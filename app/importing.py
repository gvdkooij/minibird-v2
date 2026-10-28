"""Bounded, side-effect-free parsers for contact cards and Moneybird exports."""
import base64
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from email.utils import getaddresses
import hashlib
import io
import json
from pathlib import PurePosixPath
import quopri
import re
import unicodedata
import zipfile

from defusedxml import ElementTree
from openpyxl import load_workbook

MAX_UPLOAD = 10 * 1024 * 1024
MAX_EXPANDED = 64 * 1024 * 1024
MAX_ENTRIES = 500
MAX_ROWS = 5000
NS = {'c': 'urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2',
      'a': 'urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2'}


class ImportProblem(ValueError):
    pass


def norm(value):
    return ' '.join(unicodedata.normalize('NFKC', str(value or '')).casefold().split())


def clean(value, limit=2000):
    value = str(value or '').strip()
    if len(value) > limit or any(ord(c) < 32 and c not in '\n\r\t' for c in value):
        raise ImportProblem('Een tekstveld is te lang of bevat ongeldige tekens.')
    return value


def email(value):
    value = clean(value, 254)
    if value and not re.fullmatch(r'[^\s@<>\r\n]+@[^\s@<>\r\n]+\.[^\s@<>\r\n]+', value):
        raise ImportProblem('Een contact heeft een ongeldig e-mailadres. Corrigeer het bronbestand.')
    return value


def contact(**values):
    result = {key: clean(values.get(key, ''), 500) for key in ('name','email','address','kvk','vat','contact','phone')}
    if not result['name']:
        raise ImportProblem('Een contact mist een naam.')
    raw_email = result['email']
    candidates = [address for _, address in getaddresses([raw_email.replace(';', ',')])]
    valid = []
    for address in candidates:
        try:
            if email(address): valid.append(address)
        except ImportProblem:
            pass
    result['email'] = valid[0] if valid else ''
    result['warnings'] = []
    if raw_email and (len(candidates)!=1 or len(valid)!=1 or raw_email!=result['email']):
        result['warnings'].append('E-mailadres bevat meerdere adressen of is ongeldig. Het eerste geldige adres wordt gebruikt; het origineel blijft bewaard.')
    result['archived'] = int(bool(values.get('archived', False)))
    result['external_id'] = clean(values.get('external_id',''), 100)
    result['extra'] = {**values.get('extra', {}), 'origineel_email': raw_email}
    return result


def vparts(value):
    return re.split(r'(?<!\\);', value)


def unescape(value):
    return re.sub(r'\\([nN,;\\])', lambda m: '\n' if m[1].lower() == 'n' else m[1], value)


def parse_vcards(data):
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise ImportProblem('Gebruik een UTF-8 vCard-export.') from None
    text = re.sub(r'\r?\n[ \t]', '', text)
    text = re.sub(r'=\r?\n', '', text)
    cards, fields = [], None
    for raw in text.splitlines():
        if not raw.strip():
            continue
        if raw.upper() == 'BEGIN:VCARD':
            if fields is not None: raise ImportProblem('Ongeldig begin van een vCard.')
            fields = defaultdict(list)
        elif raw.upper() == 'END:VCARD':
            if fields is None: raise ImportProblem('Ongeldig einde van een vCard.')
            get = lambda key: fields.get(key, [''])[0]
            person = get('FN') or ' '.join(unescape(p) for p in vparts(get('N'))[1:2] + vparts(get('N'))[:1]).strip()
            organization = unescape(vparts(get('ORG'))[0])
            addresses = []
            for address in fields.get('ADR', []):
                p = [unescape(x) for x in vparts(address)]
                p += [''] * (7-len(p))
                addresses.append('\n'.join(x for x in [p[0],p[1],p[2],' '.join(x for x in [p[5],p[3]] if x),p[4],p[6]] if x))
            cards.append(contact(name=organization or unescape(person), contact=unescape(person) if organization else '',
                                 email=unescape(get('EMAIL')), address=addresses[0] if addresses else '',
                                 phone=unescape(get('TEL')), extra={'vcard': dict(fields)}))
            fields = None
            if len(cards)>MAX_ROWS: raise ImportProblem('Te veel contactkaarten; splits de export.')
        elif fields is not None:
            if ':' not in raw: raise ImportProblem('Ongeldig vCard-veld.')
            key, value = raw.split(':',1)
            params = key.upper().split(';')
            name = params[0].split('.')[-1]
            if 'ENCODING=QUOTED-PRINTABLE' in params:
                try: value = quopri.decodestring(value).decode('utf-8')
                except UnicodeDecodeError: raise ImportProblem('Ongeldige vCard-codering.') from None
            fields[name].append(value)
        else:
            raise ImportProblem('Dit bestand bevat geen geldige vCards.')
    if fields is not None or not cards: raise ImportProblem('Het vCard-bestand is leeg of onvolledig.')
    return {'contacts': cards, 'documents': [], 'warnings': [], 'files': []}


def checked_zip(data):
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        entries = archive.infolist()
        if len(entries)>MAX_ENTRIES or sum(x.file_size for x in entries)>MAX_EXPANDED:
            raise ImportProblem('Het archief is te groot na uitpakken of bevat te veel bestanden.')
        seen = set()
        for entry in entries:
            path = PurePosixPath(entry.filename)
            if entry.filename in seen or path.is_absolute() or '..' in path.parts or '\\' in entry.filename or entry.flag_bits & 1 or (entry.external_attr >> 16) & 0o170000 == 0o120000:
                raise ImportProblem('Het archief bevat onveilige, dubbele of versleutelde bestandsnamen.')
            seen.add(entry.filename)
            if entry.file_size>MAX_UPLOAD or entry.file_size > max(1, entry.compress_size)*1000:
                raise ImportProblem('Een bestand in het archief is te groot of te sterk gecomprimeerd.')
        return archive
    except zipfile.BadZipFile:
        raise ImportProblem('Ongeldig of beschadigd ZIP-bestand.') from None


def spreadsheet(data, sheet=None):
    # XLSX is itself a ZIP: check its expansion before giving it to openpyxl.
    with checked_zip(data):
        pass
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        ws = workbook[sheet] if sheet else workbook.worksheets[0]
        if ws.max_row and ws.max_row>MAX_ROWS+1 or ws.max_column and ws.max_column>150:
            raise ImportProblem('Het werkblad bevat te veel rijen of kolommen.')
        iterator = ws.iter_rows(values_only=True)
        headers = next(iterator, ())
        if not headers: raise ImportProblem('Het werkblad heeft geen kolomkoppen.')
        headers = [str(x or '').strip() for x in headers]
        rows = []
        for index, values in enumerate(iterator):
            if index>=MAX_ROWS or len(values)>150: raise ImportProblem('Het werkblad is te groot.')
            if any(x is not None for x in values):
                row = {}
                for key, value in zip(headers, values):
                    # Moneybird has duplicate headers; first occurrence is authoritative.
                    if key and key not in row: row[key] = value
                rows.append(row)
        return rows
    finally:
        workbook.close()


def serial(value):
    if isinstance(value, (datetime,date)): return value.isoformat()
    return value


def money(value):
    try:
        d = Decimal(str(value))
        if not d.is_finite() or abs(d)>1000000000: raise ValueError()
        return int((d*100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    except (ValueError, InvalidOperation):
        raise ImportProblem('Een bedrag in de export is ongeldig.') from None


def decimal(value):
    try:
        text = str(value)
        d = Decimal(text)
        if len(text)>40 or not d.is_finite() or abs(d)>1000000000: raise ValueError()
        return d
    except (ValueError, InvalidOperation):
        raise ImportProblem('Een hoeveelheid of prijs is ongeldig.') from None


def day(value):
    if isinstance(value, datetime): return value.date().isoformat()
    if isinstance(value, date): return value.isoformat()
    try: return date.fromisoformat(str(value)).isoformat()
    except ValueError: raise ImportProblem('Een document heeft een ongeldige datum.') from None


def safe_number(value):
    text = str(value or '')
    if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]{0,79}', text):
        raise ImportProblem('Een documentnummer is leeg of bevat onveilige tekens.')
    return text


def node_text(node, path, default=''):
    return node.findtext(path, default=default, namespaces=NS).strip()


def party(root, role):
    node = root.find('a:'+role+'/a:Party', NS)
    if node is None: raise ImportProblem('UBL mist afzender of ontvanger.')
    addr = node.find('a:PostalAddress', NS)
    address = ''
    if addr is not None:
        address = '\n'.join(filter(None,[node_text(addr,'c:StreetName'),node_text(addr,'c:AdditionalStreetName'),' '.join(filter(None,[node_text(addr,'c:PostalZone'),node_text(addr,'c:CityName')])),node_text(addr,'a:Country/c:IdentificationCode')]))
    return contact(name=node_text(node,'a:PartyName/c:Name') or node_text(node,'a:PartyLegalEntity/c:RegistrationName'),
                   address=address, email=node_text(node,'a:Contact/c:ElectronicMail'),contact=node_text(node,'a:Contact/c:Name'),
                   kvk=node_text(node,'a:PartyLegalEntity/c:CompanyID'),vat=node_text(node,'a:PartyTaxScheme/c:CompanyID'))


def source_status(value, kind):
    mappings = {'draft':'draft','open':'sent','scheduled':'draft','pending_payment':'sent','pending':'sent','sent':'sent','late':'sent','reminded':'sent','paid':'paid','accepted':'accepted','billed':'accepted','rejected':'rejected','cancelled':'cancelled','uncollectible':'uncollectible'}
    key = str(value or '').lower()
    if key not in mappings: raise ImportProblem('Onbekende documentstatus in export: '+clean(key,50))
    result = mappings[key]
    if kind=='quote' and result in ('paid','uncollectible') or kind=='invoice' and result in ('accepted','rejected'):
        raise ImportProblem('Documenttype en bronstatus komen niet overeen.')
    return result


def ubl_document(data, summary):
    root = ElementTree.fromstring(data, forbid_dtd=True)
    if root.tag != '{urn:oasis:names:specification:ubl:schema:xsd:Invoice-2}Invoice':
        raise ImportProblem('Alleen UBL-verkoopfacturen worden ondersteund.')
    if node_text(root,'c:DocumentCurrencyCode')!='EUR': raise ImportProblem('Alleen EUR-documenten worden ondersteund.')
    if str(summary.get('valuta','EUR')).upper()!='EUR': raise ImportProblem('Alleen EUR-documenten worden ondersteund.')
    number = safe_number(node_text(root,'c:ID'))
    if number != safe_number(summary.get('factuurnummer')): raise ImportProblem('UBL-nummer komt niet overeen met het overzicht.')
    lines=[]
    for line in root.findall('a:InvoiceLine',NS):
        quantity=decimal(node_text(line,'c:InvoicedQuantity'))
        price=decimal(node_text(line,'a:Price/c:PriceAmount'))
        rate=decimal(node_text(line,'a:Item/a:ClassifiedTaxCategory/c:Percent'))
        if quantity<=0 or rate not in (0,9,21): raise ImportProblem('Niet-ondersteunde UBL-regel of btw-percentage.')
        net=money(node_text(line,'c:LineExtensionAmount'))
        vat=int((Decimal(net)*rate/100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
        lines.append(dict(description=clean(node_text(line,'a:Item/c:Description') or node_text(line,'a:Item/c:Name'),10000),quantity=str(quantity),price=str(price),rate=str(rate),net=net,vat=vat))
    if not lines or len(lines)>1000: raise ImportProblem('Ongeldig aantal UBL-regels.')
    # Moneybird may calculate VAT before rounding the line net (e.g. an inclusive
    # EUR 1,200 price). Preserve the declared tax subtotal, including that cent.
    tax_groups = root.findall('a:TaxTotal/a:TaxSubtotal', NS)
    seen_rates = set()
    for group in tax_groups:
        rate = decimal(node_text(group,'a:TaxCategory/c:Percent'))
        if rate in seen_rates: raise ImportProblem('Meerdere UBL-btw-groepen met hetzelfde tarief worden niet ondersteund.')
        seen_rates.add(rate)
        group_lines = [line for line in lines if Decimal(line['rate'])==rate]
        target = money(node_text(group,'c:TaxAmount'))
        delta = target - sum(line['vat'] for line in group_lines)
        if not group_lines or money(node_text(group,'c:TaxableAmount'))!=sum(line['net'] for line in group_lines) or abs(delta)>len(group_lines):
            raise ImportProblem('De UBL-btw-groepen sluiten niet aan op de regels.')
        for index in range(abs(delta)):
            group_lines[index % len(group_lines)]['vat'] += 1 if delta>0 else -1
    if set(Decimal(line['rate']) for line in lines)!=seen_rates or sum(line['vat'] for line in lines)!=money(node_text(root,'a:TaxTotal/c:TaxAmount')):
        raise ImportProblem('Het UBL-btw-totaal sluit niet aan.')
    total=money(node_text(root,'a:LegalMonetaryTotal/c:TaxInclusiveAmount'))
    if sum(l['net'] for l in lines)!=money(node_text(root,'a:LegalMonetaryTotal/c:TaxExclusiveAmount')) or sum(l['net']+l['vat'] for l in lines)!=total or money(summary.get('totaalprijs inclusief btw'))!=total:
        raise ImportProblem('UBL-regels, btw en documenttotaal sluiten niet op elkaar aan. Import is gestopt.')
    company=party(root,'AccountingSupplierParty')
    company['iban']=node_text(root,'a:PaymentMeans/a:PayeeFinancialAccount/c:ID')
    customer=party(root,'AccountingCustomerParty')
    return dict(kind='invoice',number=number,date=day(node_text(root,'c:IssueDate')),due=day(node_text(root,'c:DueDate')),customer=customer,company=company,lines=lines,total=total,status=source_status(summary.get('status'),'invoice'),original_status=str(summary.get('status')),external_id=clean(summary.get('id'),100))


def parse_moneybird(data):
    with checked_zip(data) as archive:
        names=archive.namelist()
        contact_files=[n for n in names if n.startswith('contacts/') and n.endswith('.xlsx')]
        if len(contact_files)!=1: raise ImportProblem('Geen eenduidige Moneybird-contactenexport gevonden in het ZIP-bestand.')
        rows=spreadsheet(archive.read(contact_files[0]),'Contacten')
        contacts=[]; by_number={}
        contact_people=spreadsheet(archive.read(contact_files[0]),'Contactpersonen')
        for row in rows:
            person=' '.join(str(row.get(k) or '') for k in ('firstname','lastname')).strip()
            address='\n'.join(filter(None,[str(row.get('address1') or ''),str(row.get('address2') or ''),' '.join(str(row.get(k) or '') for k in ('zipcode','city')).strip(),str(row.get('country') or '')]))
            c=contact(name=row.get('company_name') or person,contact=row.get('attention') or person,email=row.get('email'),phone=row.get('phone'),address=address,kvk=row.get('chamber_of_commerce'),vat=row.get('tax_number'),archived=str(row.get('archived')).lower()=='true',external_id=row.get('id'),extra={'moneybird':{k:serial(v) for k,v in row.items()},'contactpersonen':[{k:serial(v) for k,v in p.items()} for p in contact_people if str(p.get('customer_id'))==str(row.get('customer_id'))]})
            contacts.append(c); by_number[str(row.get('customer_id'))]=c
        documents=[]; used={contact_files[0]}; company=None
        invoice_files=[n for n in names if n.startswith('sales_invoices/') and n.endswith('.xlsx')]
        if len(invoice_files)>1: raise ImportProblem('Meerdere factuuroverzichten gevonden.')
        for name in invoice_files:
            used.add(name)
            for summary in spreadsheet(archive.read(name)):
                number=safe_number(summary.get('factuurnummer'))
                xml='sales_invoices_ubl_archive/'+number+'.xml'
                if xml not in names: raise ImportProblem('UBL ontbreekt voor een factuur; deze export kan niet volledig worden geïmporteerd.')
                d=ubl_document(archive.read(xml),summary); used.add(xml)
                d['contact_key']=str(summary.get('klantnummer'))
                if d['contact_key'] not in by_number: raise ImportProblem('Een factuur verwijst naar een ontbrekend contact.')
                if company is not None and norm(company['name'])!=norm(d['company']['name']): raise ImportProblem('Export bevat meerdere afzenders.')
                company=d['company']
                pdf_path='sales_invoices_archive/'+d['date'][:4]+'/'+number+'.pdf'
                d['pdf']=read_pdf(archive,pdf_path) if pdf_path in names else embedded_pdf(archive.read(xml))
                if pdf_path in names: used.add(pdf_path)
                documents.append(d)
        quote_files=[n for n in names if n.startswith('estimates/') and n.endswith('.xlsx')]
        quote_detail_files=[n for n in names if n.startswith('estimates_details/') and n.endswith('.xlsx')]
        if len(quote_files)>1 or len(quote_detail_files)>1: raise ImportProblem('Meerdere offerteoverzichten gevonden.')
        for name in quote_files:
            used.add(name)
            summaries=spreadsheet(archive.read(name))
            if not summaries: continue
            if not company: raise ImportProblem('Afzendergegevens ontbreken: voeg UBL-facturen toe aan de jaarexport.')
            if not quote_detail_files: raise ImportProblem('Offerte-regels ontbreken in deze export.')
            used.update(quote_detail_files)
            detail=spreadsheet(archive.read(quote_detail_files[0]))
            for summary in summaries:
                if str(summary.get('valuta','')).upper()!='EUR': raise ImportProblem('Alleen EUR-documenten worden ondersteund.')
                number=safe_number(summary.get('offertenummer')); lines=[]
                for row in detail:
                    if str(row.get('offertenummer'))!=number: continue
                    rate_text=str(row.get('btw') or '')
                    match=re.fullmatch(r'(0|9|21)% btw',rate_text)
                    if not match: raise ImportProblem('Niet-ondersteund btw-tarief op een offerte.')
                    rate=int(match[1]); net=money(row.get('totaalprijs exclusief btw inclusief korting'))
                    quantity=decimal(row.get('aantal (decimaal)')); price=decimal(row.get('prijs'))
                    if quantity<=0: raise ImportProblem('Ongeldige hoeveelheid in offerte.')
                    lines.append(dict(description=clean(row.get('omschrijving'),10000),quantity=str(quantity),price=str(price),rate=str(rate),net=net,vat=int((Decimal(net)*rate/100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))))
                total=money(summary.get('totaalprijs inclusief btw'))
                if not lines or sum(l['net'] for l in lines)!=money(summary.get('totaalprijs exclusief btw')) or sum(l['net']+l['vat'] for l in lines)!=total:
                    raise ImportProblem('Offerte-regels, btw en totaal sluiten niet op elkaar aan.')
                key=str(summary.get('klantnummer'))
                if key not in by_number: raise ImportProblem('Offerte verwijst naar een ontbrekend contact.')
                issued=day(summary.get('offerte datum'))
                pdf_path='estimates_archive/'+issued[:4]+'/'+number+'.pdf'
                if pdf_path not in names: raise ImportProblem('De originele offerte-PDF ontbreekt.')
                used.add(pdf_path)
                documents.append(dict(kind='quote',number=number,date=issued,due=day(summary.get('vervaldatum')),customer=by_number[key],company=company,lines=lines,total=total,status=source_status(summary.get('status'),'quote'),original_status=str(summary.get('status')),external_id=clean(summary.get('id'),100),contact_key=key,pdf=read_pdf(archive,pdf_path)))
        for c in contacts:
            c['customer_number']=next((key for key,value in by_number.items() if value is c),'')
        files=[{'name':n,'imported':n in used} for n in names if not n.endswith('/')]
        warnings=['Aanvullende rapporten, producten, periodieke facturen, transacties en auditfile worden alleen in het bronarchief bewaard; er worden geen boekingen van gemaakt.', 'Geïmporteerde documenten zijn alleen-lezen. Bestaande klanten worden niet overschreven. Extra contactvelden en contactpersonen blijven bij het importcontact bewaard.']
        if any(not f['imported'] for f in files): warnings.append('Het volledige oorspronkelijke ZIP-bestand blijft na bevestiging als download beschikbaar.')
        return {'contacts':contacts,'documents':documents,'warnings':warnings,'files':files}


def read_pdf(archive,path):
    data=archive.read(path)
    if not data.startswith(b'%PDF-'): raise ImportProblem('Een origineel document is geen PDF.')
    return base64.b64encode(data).decode('ascii')


def embedded_pdf(xml):
    root=ElementTree.fromstring(xml,forbid_dtd=True)
    for node in root.findall('a:AdditionalDocumentReference/a:Attachment/c:EmbeddedDocumentBinaryObject',NS):
        if node.get('mimeCode')=='application/pdf':
            try: data=base64.b64decode(''.join((node.text or '').split()),validate=True)
            except ValueError: raise ImportProblem('Ongeldige ingebedde PDF.') from None
            if data.startswith(b'%PDF-') and len(data)<=MAX_UPLOAD: return base64.b64encode(data).decode('ascii')
    raise ImportProblem('De originele factuur-PDF ontbreekt.')


def parse_upload(filename,data):
    if not data or len(data)>MAX_UPLOAD: raise ImportProblem('Upload een niet-leeg bestand van maximaal 10 MB.')
    try:
        extension=filename.lower().rsplit('.',1)[-1]
        if extension=='vcf': result=parse_vcards(data)
        elif extension=='zip': result=parse_moneybird(data)
        else: raise ImportProblem('Kies een .vcf- of Moneybird .zip-bestand.')
        # Bound the stored preview even when source fields contain repeated data.
        if len(json.dumps(result,ensure_ascii=False))>MAX_EXPANDED:
            raise ImportProblem('Te veel gegevens voor één import.')
        return result
    except ImportProblem:
        raise
    except Exception:
        raise ImportProblem('Het bestand is beschadigd of heeft geen ondersteunde exportindeling.') from None


def fingerprint(document):
    # Original PDF included, so a changed source cannot silently replace history.
    return hashlib.sha256(json.dumps(document,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
